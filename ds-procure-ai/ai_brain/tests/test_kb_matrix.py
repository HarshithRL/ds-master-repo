"""W8: matrix/deck from SQL evidence + pricing_line; distinct render states."""

from __future__ import annotations

import shutil
import tempfile
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ai_brain.core.procurement_ai.insights import build_insight_payload
from ai_brain.core.procurement_ai.packs.compare_xlsx import build_comparison_xlsx
from ai_brain.core.procurement_ai.packs.matrix_from_evidence import (
    pick_evidence_row,
    render_evidence_cell,
)
from ai_brain.core.procurement_ai.packs.matrix_layout import (
    EXTRACTION_FAILED,
    NOT_OFFERED,
    NOT_STATED,
    REVIEWED_MARK,
    SOURCE_QUALITY_PREFIX,
)
from ai_brain.core.procurement_ai.packs.steerco_ppt import build_steerco_ppt
from shared.db.models import Evidence, QualityFlag, utc_now_iso


@pytest.fixture()
def db_env(monkeypatch):
    from shared.db import init_db
    from shared.db import connection as db_connection

    if db_connection._engine is not None:
        db_connection._engine.dispose()
    db_connection._engine = None
    db_connection._SessionLocal = None
    tmp = tempfile.mkdtemp()
    db_path = Path(tmp) / "test.sqlite"
    monkeypatch.setenv("MATE_SQLITE_PATH", str(db_path))
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(Path(tmp) / "projects"))
    init_db()
    yield db_path
    if db_connection._engine is not None:
        db_connection._engine.dispose()
    db_connection._engine = None
    db_connection._SessionLocal = None
    shutil.rmtree(tmp, ignore_errors=True)


def _seed_warm_kb():
    from shared.db.connection import session_scope
    from shared.db.models import (
        AppUser,
        Artifact,
        PricingLine,
        Project,
        Requirement,
        Vendor,
    )

    now = utc_now_iso()
    with session_scope() as session:
        session.add(AppUser(workspace_id="u1", created_at=now, updated_at=now))
        session.add(
            Project(
                id="p1",
                code="PRJ-TEST",
                owner_id="u1",
                module_id="procurement-ai-assistant",
                name="SWIFT CSP assessment",
                business_process="Indirect",
                category="SWIFT CSCF",
                status="Active",
                priority="Medium",
                progress=0,
                created_at=now,
                updated_at=now,
            )
        )
        session.add(
            Artifact(
                id="art-ey",
                project_id="p1",
                uploaded_by="u1",
                kind="upload",
                original_name="EY_proposal.pdf",
                storage_relpath="rel/ey",
                created_at=now,
                parse_status="ok",
                vendor_id="ey",
                doc_role="proposal",
            )
        )
        session.add(
            Artifact(
                id="art-bdo",
                project_id="p1",
                uploaded_by="u1",
                kind="upload",
                original_name="Etex_Group_security_proposal.pdf",
                storage_relpath="rel/bdo",
                created_at=now,
                parse_status="ok",
                vendor_id="bdo",
                doc_role="proposal",
            )
        )
        session.add(
            Artifact(
                id="art-csr",
                project_id="p1",
                uploaded_by="u1",
                kind="upload",
                original_name="KPMG_CSR.pdf",
                storage_relpath="rel/csr",
                created_at=now,
                parse_status="ok",
                vendor_id="ey",
                doc_role="csr",
            )
        )
        session.add(
            Artifact(
                id="art-rfp",
                project_id="p1",
                uploaded_by="u1",
                kind="upload",
                original_name="Etex_RFP_Swift CSP.pdf",
                storage_relpath="rel/rfp",
                created_at=now,
                parse_status="ok",
                doc_role="rfp",
            )
        )
        session.add(
            Vendor(
                id="p1:ey",
                project_id="p1",
                name="EY",
                legal_entity="EY Bedrijfsrevisoren BV",
                created_at=now,
            )
        )
        session.add(
            Vendor(
                id="p1:bdo",
                project_id="p1",
                name="BDO",
                legal_entity="BDO Advisory BV",
                created_at=now,
            )
        )
        session.add(
            Requirement(
                id="req-ns",
                project_id="p1",
                ordinal=1,
                section="scope",
                text="Two BIC codes in scope",
                category="scope",
                mandatory=1,
                created_at=now,
            )
        )
        session.add(
            Requirement(
                id="req-no",
                project_id="p1",
                ordinal=2,
                section="commercial",
                text="Dual pricing T&M and Fixed Price",
                category="commercial",
                mandatory=1,
                created_at=now,
            )
        )
        session.add(
            Requirement(
                id="req-fail",
                project_id="p1",
                ordinal=3,
                section="deliverable",
                text="CSR / EcoVadis disclosure",
                category="deliverable",
                mandatory=0,
                created_at=now,
            )
        )
        session.add(
            Requirement(
                id="req-met",
                project_id="p1",
                ordinal=4,
                section="scope",
                text="CSCF 2026 Architecture A4",
                category="scope",
                mandatory=1,
                created_at=now,
            )
        )
        session.add(
            Evidence(
                id="ev-ns",
                project_id="p1",
                requirement_id="req-ns",
                vendor_id="bdo",
                status="not_stated",
                model_version="v1",
                prompt_version="p1",
                extracted_at=now,
            )
        )
        session.add(
            Evidence(
                id="ev-no",
                project_id="p1",
                requirement_id="req-no",
                vendor_id="bdo",
                status="not_offered",
                claim_text="We do not offer dual pricing",
                artifact_id="art-bdo",
                page=2,
                confidence=0.9,
                model_version="v1",
                prompt_version="p1",
                extracted_at=now,
            )
        )
        session.add(
            Evidence(
                id="ev-fail",
                project_id="p1",
                requirement_id="req-fail",
                vendor_id="bdo",
                status="extraction_failed",
                claim_text="HALLUCINATED ANSWER that must never render",
                artifact_id="art-bdo",
                page=9,
                confidence=0.1,
                model_version="v1",
                prompt_version="p1",
                extracted_at=now,
            )
        )
        session.add(
            Evidence(
                id="ev-met-old",
                project_id="p1",
                requirement_id="req-met",
                vendor_id="bdo",
                status="not_stated",
                model_version="v1",
                prompt_version="p1",
                extracted_at="2026-01-01T00:00:00+00:00",
                reviewed_by="analyst@etex",
                reviewed_at="2026-01-02T00:00:00+00:00",
            )
        )
        session.add(
            Evidence(
                id="ev-met-new",
                project_id="p1",
                requirement_id="req-met",
                vendor_id="bdo",
                status="met",
                claim_text="we confirm CSCF 2026 Architecture A4",
                artifact_id="art-csr",
                page=1,
                confidence=0.95,
                model_version="v2",
                prompt_version="p1",
                extracted_at="2026-09-01T00:00:00+00:00",
            )
        )
        session.add(
            Evidence(
                id="ev-ey-met",
                project_id="p1",
                requirement_id="req-met",
                vendor_id="ey",
                status="met",
                claim_text="EY will assess CSCF 2026",
                artifact_id="art-csr",
                page=3,
                confidence=0.8,
                model_version="v1",
                prompt_version="p1",
                extracted_at=now,
            )
        )
        session.add(
            QualityFlag(
                id="qf-csr",
                project_id="p1",
                artifact_id="art-csr",
                kind="low_yield",
                detail="747 tokens from a 418 KB file",
                severity="high",
                created_at=now,
            )
        )
        session.add(
            PricingLine(
                id="pl-bdo",
                project_id="p1",
                vendor_id="bdo",
                offer_type="time_and_materials",
                role="manager",
                rate=1200,
                rate_unit="day",
                currency="EUR",
                quantity=40,
                line_total=48000,
                etex_effort_days=5,
                artifact_id="art-bdo",
                page=4,
                discount_terms="T&M estimate, incl. 10% discount. No fixed-price offer.",
                verified=1,
                model_version="v1",
            )
        )


def _project_insights():
    class _Project:
        id = "p1"
        code = "PRJ-TEST"
        name = "SWIFT CSP assessment"
        business_process = "Indirect"
        category = "SWIFT CSCF"
        description = "Professional services days"

    class _Artifact:
        id = "art-ey"
        original_name = "EY_proposal.pdf"
        parse_status = "ok"

    return build_insight_payload(_Project(), [_Artifact()], [])


def _req_sheet(path: str):
    return load_workbook(path)["Procurement Requirements"]


_LEAD = {"ID", "Requirement", "Category", "Mandatory", "Target / Detail"}


def _headers(ws) -> list[str | None]:
    return [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]


def _vendor_cols(headers: list[str | None]) -> list[tuple[int, str]]:
    return [
        (index, str(name))
        for index, name in enumerate(headers, start=1)
        if name and name not in _LEAD
    ]


def _row_for(ws, requirement_id: str) -> int:
    for row in range(5, (ws.max_row or 5) + 1):
        if ws.cell(row, 1).value == requirement_id:
            return row
    raise AssertionError(f"requirement {requirement_id} not found")


def _cell_map(path: str) -> dict[tuple[str, str], str]:
    ws = _req_sheet(path)
    headers = _headers(ws)
    vendors = _vendor_cols(headers)
    out: dict[tuple[str, str], str] = {}
    for row in range(5, (ws.max_row or 5) + 1):
        rid = ws.cell(row, 1).value
        if not rid or rid in {item[1] for item in vendors}:
            continue
        if str(rid) in {"scope", "commercial", "deliverable", "timeline", "legal"}:
            continue
        for col, vendor in vendors:
            out[(str(rid), vendor)] = str(ws.cell(row, col).value or "")
    return out


def _fill_rgb(cell) -> str:
    color = cell.fill.fgColor
    rgb = getattr(color, "rgb", None)
    return str(rgb)[-6:].upper() if rgb else ""


def test_render_states_are_textually_and_visually_distinct():
    ns = render_evidence_cell(
        Evidence(
            id="a",
            project_id="p",
            requirement_id="r",
            vendor_id="v",
            status="not_stated",
            model_version="v",
            prompt_version="p",
            extracted_at="t",
        )
    )
    offered = render_evidence_cell(
        Evidence(
            id="b",
            project_id="p",
            requirement_id="r",
            vendor_id="v",
            status="not_offered",
            claim_text="no dual pricing",
            artifact_id="art",
            page=1,
            model_version="v",
            prompt_version="p",
            extracted_at="t",
        )
    )
    failed = render_evidence_cell(
        Evidence(
            id="c",
            project_id="p",
            requirement_id="r",
            vendor_id="v",
            status="extraction_failed",
            claim_text="HALLUCINATED ANSWER that must never render",
            artifact_id="art",
            page=1,
            model_version="v",
            prompt_version="p",
            extracted_at="t",
        )
    )
    assert ns.text.split("\n")[0] == NOT_STATED
    assert offered.text.split("\n")[0] == NOT_OFFERED
    assert failed.text.split("\n")[0] == EXTRACTION_FAILED
    assert ns.text.split("\n")[0] != offered.text.split("\n")[0] != failed.text.split("\n")[0]
    assert "HALLUCINATED" not in failed.text
    assert "[p.1]" in offered.text
    assert offered.citation
    assert "p.1" in (offered.citation or "")
    assert {ns.status, offered.status, failed.status} == {
        "not_stated",
        "not_offered",
        "extraction_failed",
    }


def test_pick_evidence_prefers_reviewed_row():
    old = Evidence(
        id="old",
        project_id="p",
        requirement_id="r",
        vendor_id="v",
        status="not_stated",
        model_version="v1",
        prompt_version="p",
        extracted_at="2026-01-01T00:00:00+00:00",
        reviewed_by="analyst@etex",
        reviewed_at="2026-01-02T00:00:00+00:00",
    )
    new = Evidence(
        id="new",
        project_id="p",
        requirement_id="r",
        vendor_id="v",
        status="met",
        claim_text="newer unreviewed claim",
        artifact_id="art",
        page=1,
        model_version="v2",
        prompt_version="p",
        extracted_at="2026-09-01T00:00:00+00:00",
    )
    chosen = pick_evidence_row([new, old])
    assert chosen.id == "old"
    rendered = render_evidence_cell(chosen)
    assert NOT_STATED in rendered.text
    assert REVIEWED_MARK in rendered.text
    assert "newer unreviewed" not in rendered.text


def test_high_severity_flag_adds_source_quality_warning():
    row = Evidence(
        id="e",
        project_id="p",
        requirement_id="r",
        vendor_id="v",
        status="met",
        claim_text="EY will assess CSCF 2026",
        artifact_id="art-csr",
        page=3,
        confidence=0.8,
        model_version="v",
        prompt_version="p",
        extracted_at="t",
    )
    flag = QualityFlag(
        id="qf",
        project_id="p",
        artifact_id="art-csr",
        kind="low_yield",
        severity="high",
        created_at="t",
    )
    rendered = render_evidence_cell(row, [flag], artifact_names={"art-csr": "KPMG_CSR.pdf"})
    assert SOURCE_QUALITY_PREFIX in rendered.text
    assert "low_yield" in rendered.text
    assert "EY will assess CSCF 2026" in rendered.text
    assert "[p.3]" in rendered.text
    assert rendered.citation and "KPMG_CSR.pdf" in rendered.citation


def test_matrix_regenerates_from_warm_evidence(db_env, monkeypatch):
    _seed_warm_kb()
    insights = _project_insights()

    def _boom(*_args, **_kwargs):
        raise AssertionError("LLM must not run in the xlsx file-writing path")

    monkeypatch.setattr(
        "ai_brain.core.procurement_ai.packs.matrix_extract.extract_matrix",
        _boom,
    )
    started = time.perf_counter()
    first = build_comparison_xlsx("p1", insights, thread_id="t-w8")
    elapsed = time.perf_counter() - started
    assert first["status"] == "ready"
    assert first["llm_used"] is False
    assert first["evidence_used"] is True
    assert elapsed < 10
    assert Path(first["path"]).exists()

    cells = _cell_map(first["path"])
    assert ("req-ns", "BDO") in cells
    assert cells[("req-ns", "BDO")].startswith(NOT_STATED)
    assert cells[("req-no", "BDO")].startswith(NOT_OFFERED)
    failed_text = cells[("req-fail", "BDO")]
    assert failed_text.startswith(EXTRACTION_FAILED)
    assert "HALLUCINATED" not in failed_text
    pinned = cells[("req-met", "BDO")]
    assert NOT_STATED in pinned
    assert REVIEWED_MARK in pinned
    assert "we confirm CSCF" not in pinned
    ey_met = cells[("req-met", "EY")]
    assert "EY will assess CSCF 2026" in ey_met
    assert SOURCE_QUALITY_PREFIX in ey_met
    assert "[p.3]" in ey_met

    ws = _req_sheet(first["path"])
    headers = _headers(ws)
    assert "Mandatory" in headers
    assert "BDO" in headers
    assert "EY" in headers
    assert "Etex" not in headers
    assert any(
        str(ws.cell(row, 1).value).lower() == "scope"
        for row in range(5, (ws.max_row or 5) + 1)
        if ws.cell(row, 1).value
    )
    ns_row = _row_for(ws, "req-ns")
    assert ws.cell(ns_row, 4).value == "Yes"
    bdo_col = headers.index("BDO") + 1
    fills = {
        _fill_rgb(ws.cell(_row_for(ws, "req-ns"), bdo_col)),
        _fill_rgb(ws.cell(_row_for(ws, "req-no"), bdo_col)),
        _fill_rgb(ws.cell(_row_for(ws, "req-fail"), bdo_col)),
    }
    assert "D9D9D9" in fills
    assert "F4CCCC" in fills
    assert "F4B183" in fills
    assert len(fills) == 3
    ey_met_cell = ws.cell(_row_for(ws, "req-met"), headers.index("EY") + 1)
    assert ey_met_cell.comment is not None
    assert "p.3" in ey_met_cell.comment.text
    assert "KPMG_CSR.pdf" in ey_met_cell.comment.text

    wb = load_workbook(first["path"])
    assert "Pricing lines" in wb.sheetnames
    pricing = wb["Pricing lines"]
    assert pricing["A2"].value == "BDO"
    assert pricing["B2"].value == "time_and_materials"
    bafo = wb["Final BAFO Comparison"]
    bdo_bafo = 3
    for col in range(3, bafo.max_column + 1):
        if bafo.cell(1, col).value == "BDO":
            bdo_bafo = col
            break
    assert "EUR 48000" in str(bafo.cell(4, bdo_bafo).value)
    assert bafo.cell(4, bdo_bafo).comment is not None
    assert "p.4" in bafo.cell(4, bdo_bafo).comment.text
    assert "BDO Advisory BV" in str(bafo.cell(2, bdo_bafo).value)
    assert "no llm" in str(wb["Cover"]["A2"].value).lower()
    sources = wb["Sources & Notes"]
    assert r"D:\Work" not in str(sources["A1"].value)
    blob = " ".join(str(sources.cell(r, 2).value or "") for r in range(1, 12))
    assert "EY_proposal.pdf" in blob
    assert "Etex_RFP_Swift CSP.pdf" in blob

    second = build_comparison_xlsx("p1", insights, thread_id="t-w8-again")
    assert _cell_map(second["path"]) == cells

    ppt = build_steerco_ppt("p1", thread_id="t-w8-ppt")
    assert ppt["status"] == "ready"
    from pptx import Presentation

    deck = Presentation(ppt["path"])
    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    blob = "\n".join(texts)
    assert "Matrix from SQL evidence" in blob
    assert "BDO" in blob
    assert "no LLM" in blob
