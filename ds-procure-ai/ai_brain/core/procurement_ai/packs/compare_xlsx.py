"""Comparison workbook from a request-time ComparisonMatrix. No LLM in the writer."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from ai_brain.core.procurement_ai.ontology import ComparisonMatrix
from ai_brain.core.procurement_ai.packs import store
from ai_brain.core.procurement_ai.packs.matrix_direct import write_direct
from ai_brain.core.procurement_ai.packs.matrix_extract import empty_direct, empty_indirect
from ai_brain.core.procurement_ai.packs.matrix_from_ontology import (
    direct_from_ontology,
    indirect_from_ontology,
    matrix_is_rich,
)
from ai_brain.core.procurement_ai.packs.matrix_indirect import write_indirect
from ai_brain.core.procurement_ai.packs.matrix_schema import DirectMatrix, IndirectMatrix
from ai_brain.core.procurement_ai.packs.xlsx_style import COVER_FIELDS
from ai_brain.core.procurement_ai.process_type import DIRECT_TG
from shared.logger_global import get_logger

log = get_logger(__name__, service="ai_brain")


def build_comparison_xlsx(
    project_id: str,
    insights: dict[str, Any],
    *,
    thread_id: str = "",
    chunks: list[Any] | None = None,
    artifacts: list[Any] | None = None,
    matrix: ComparisonMatrix | None = None,
) -> dict[str, Any]:
    del chunks, artifacts
    store.write_status(project_id, "xlsx", status="running", thread_id=thread_id or None)
    filename = "comparison_matrix.xlsx"
    path = store.pack_file(project_id, filename)
    process = (matrix.process_type if matrix else "") or insights.get("process_type") or ""
    rich = matrix_is_rich(matrix)
    if matrix is not None:
        insights = _insights_from_matrix(insights, matrix)
        indirect = indirect_from_ontology(matrix, insights)
        direct = direct_from_ontology(matrix, insights) if process == DIRECT_TG else None
    elif process == DIRECT_TG:
        indirect = empty_indirect(insights)
        direct = empty_direct(insights)
    else:
        indirect = empty_indirect(insights)
        direct = None
    if not isinstance(indirect, IndirectMatrix):
        indirect = empty_indirect(insights)
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    _write_cover(cover, insights, evidence_used=rich)
    _write_vendors(wb.create_sheet("Vendors"), insights)
    write_indirect(wb, insights, indirect)
    if process == DIRECT_TG:
        if not isinstance(direct, DirectMatrix):
            direct = empty_direct(insights)
        write_direct(wb, insights, direct)
    wb.save(path)
    href = store.flask_href(project_id, "xlsx")
    placeholder = not (insights.get("vendors") or [])
    pack_status = "ready" if rich else "template"
    pack_error = None
    if not process:
        pack_error = "process_type unset. Cells are ?, not N/A, until Direct or Indirect is known."
        pack_status = "template"
    elif not rich:
        pack_error = (
            "Template only. No named vendors in the comparison payload. Unevidenced cells are ?."
            if placeholder
            else "Template only — comparison agent did not return cited cells."
        )
    store.write_status(
        project_id,
        "xlsx",
        status=pack_status,
        href=href,
        filename=filename,
        thread_id=thread_id or None,
        error=pack_error,
    )
    log.info(
        "comparison xlsx {} project_id={} process={} rich={} sheets={}",
        pack_status,
        project_id,
        process or "unset",
        rich,
        wb.sheetnames,
    )
    return {
        "status": pack_status,
        "href": href,
        "filename": filename,
        "path": str(path),
        "llm_used": False,
        "evidence_used": rich,
        "sheets": list(wb.sheetnames),
        "error": pack_error,
    }


def _insights_from_matrix(insights: dict[str, Any], matrix: ComparisonMatrix) -> dict[str, Any]:
    out = dict(insights)
    out["process_type"] = matrix.process_type or out.get("process_type") or ""
    vendors = []
    for row in matrix.vendors:
        files = (row.source_files or "").split(";") if row.source_files else []
        files = [part.strip() for part in files if part.strip()]
        vendors.append(
            {
                "name": row.legal_name or row.vendor_id,
                "bid_id": row.vendor_id,
                "role": "bidder",
                "artifact_count": len(files),
                "coverage_pct": 0,
                "headline": row.source_files or "",
                "evidence": [],
            }
        )
    if vendors:
        out["vendors"] = vendors
    decision = dict(out.get("decision") or {})
    if matrix.summary:
        decision["summary"] = matrix.summary
    out["decision"] = decision
    return out


def _write_cover(ws, insights: dict[str, Any], *, evidence_used: bool) -> None:
    ws["A1"] = "Mate comparison pack"
    ws["A1"].font = Font(bold=True, size=16, color="FF6900")
    origin = (
        "Cells filled from request-time comparison agent (cited chunks). No invented numbers."
        if evidence_used
        else "Template only — no cited comparison payload."
    )
    ws["A2"] = "Draft workbook. Humans award. Empty is missing, never zero. " + origin
    ws["A4"] = "Named field"
    ws["B4"] = "Value"
    ws["A4"].font = Font(bold=True)
    ws["B4"].font = Font(bold=True)
    for row, (field, key) in enumerate(COVER_FIELDS, start=5):
        ws[f"A{row}"] = field
        ws[f"B{row}"] = insights.get(key, "")
    ws["A12"] = "FIELD_DECISION_SUMMARY"
    ws["B12"] = (insights.get("decision") or {}).get("summary") or ""
    ws["A13"] = "FIELD_WEIGHTS_LOCKED"
    ws["B13"] = "NO — draft scores only"
    ws["A14"] = "FIELD_AWARD"
    ws["B14"] = "Not awarded"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 72


def _write_vendors(ws, insights: dict[str, Any]) -> None:
    from ai_brain.core.procurement_ai.packs.xlsx_style import autosize, header_row

    header_row(ws, ["Vendor", "Role", "Files", "Coverage %", "Headline", "Evidence locator"])
    vendors = insights.get("vendors") or []
    if not vendors:
        ws["A2"] = "No vendor column yet"
        return
    for index, vendor in enumerate(vendors, start=2):
        evidence = vendor.get("evidence") or []
        ws[f"A{index}"] = vendor.get("name")
        ws[f"B{index}"] = vendor.get("role")
        ws[f"C{index}"] = vendor.get("artifact_count")
        ws[f"D{index}"] = vendor.get("coverage_pct")
        ws[f"E{index}"] = vendor.get("headline")
        ws[f"F{index}"] = evidence[0].get("locator") if evidence else "missing"
    autosize(ws)


def existing_xlsx_path(project_id: str) -> Path | None:
    status = store.read_status(project_id).get("xlsx") or {}
    filename = status.get("filename")
    if not filename:
        return None
    path = store.pack_file(project_id, filename)
    return path if path.exists() else None
