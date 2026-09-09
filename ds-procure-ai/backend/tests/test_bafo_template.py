"""Shared BAFO template: N/A on rows that do not apply. No invented numbers."""

from openpyxl import load_workbook

from ai_brain.core.procurement_ai.packs.bafo_from_template import write_bafo_workbook


def _values(path, process_type):
    write_bafo_workbook(path, {"process_type": process_type}, facts={})
    wb = load_workbook(path)
    ws = wb["Final BAFO Comparison"]
    found = {}
    for row in ws.iter_rows(min_row=3, max_col=3, values_only=False):
        label = row[0].value
        if label:
            found[str(label)] = row[2].value
    return found


def test_direct_marks_indirect_rows_na(tmp_path):
    values = _values(tmp_path / "direct.xlsx", "direct_tg")
    assert values["Indexation: (applies only to recurring components unless explicitly stated otherwise)"] == "N/A"
    assert values["Architecture Assessment"] == "N/A"
    assert values["Supply type location"] == "?"
    assert values["Score Card score"] == "To be scored by Etex evaluation team"
    assert values["Total TCO year 1"] == "?"


def test_services_marks_direct_and_software_rows_na(tmp_path):
    values = _values(tmp_path / "services.xlsx", "indirect_it_services")
    assert values["Supply type location"] == "N/A"
    assert values["Solution costs (Software, Materials)"] == "N/A"
    assert values["Service/Implementation Costs"] == "?"
    assert load_workbook(tmp_path / "services.xlsx")["Final BAFO Comparison"]["C1"].value == "Vendor 1"
