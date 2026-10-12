import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from pyspark.sql import SparkSession


ALLOWED_CATALOGS = {"automation_d", "automation_a", "automation_p"}
CONFIGURATION_FILES = {
    "general_data.xlsx": {
        "General_Data_Fields",
        "Company_Sales_Org",
        "Tax_Details",
    },
    "customer_form_rules.xlsx": {"Form_Fields", "Field_Options"},
    "credit_management_rules.xlsx": {"Credit_Mapping", "Form_Fields"},
    "assignment_rules.xlsx": {"Members", "Assignment_Pool", "Notification_Messages"},
}
PRODUCTION_BLOCKERS = ("test only", "dummy configuration", "used only to verify")
CONFIGURATION_TABLE_SCHEMA = """
source_file STRING,
sheet_name STRING,
records_json STRING,
source_sha256 STRING,
loaded_at TIMESTAMP
"""


def validate_production_configuration(catalog: str, workbook_values: list[str]) -> None:
    if catalog == "automation_d":
        return

    if any(
        marker in value.casefold()
        for value in workbook_values
        for marker in PRODUCTION_BLOCKERS
    ):
        raise ValueError(
            "ACC/PRD configuration seeding is blocked because the supplied workbooks "
            "contain test-only or dummy configuration. Replace them with approved "
            "production rules before running this job."
        )


def _cell_text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()
    return str(value)


def load_configuration_snapshot(
    config_dir: str | Path,
    catalog: str,
) -> tuple[list[tuple[str, str, str, str, datetime]], int]:
    from openpyxl import load_workbook

    root = Path(config_dir)
    loaded_at = datetime.now(timezone.utc)
    snapshot = []
    workbook_values = []
    blocker_locations = []
    row_count = 0

    for file_name, required_sheets in CONFIGURATION_FILES.items():
        file_path = root / file_name
        if file_name == "general_data.xlsx" and not file_path.is_file():
            file_path = root / "customer_reference_library.xlsx"
        if not file_path.is_file():
            raise FileNotFoundError(f"Required configuration workbook is missing: {file_path}")
        file_name = file_path.name

        source_sha256 = hashlib.sha256(file_path.read_bytes()).hexdigest()
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        missing_sheets = required_sheets.difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                f"{file_name} is missing required sheets: {', '.join(sorted(missing_sheets))}"
            )

        for worksheet in workbook.worksheets:
            iterator = worksheet.iter_rows(values_only=True)
            header = next(iterator, None)
            records = []
            if header:
                columns = [
                    _cell_text(value).strip() or f"Unnamed: {index}"
                    for index, value in enumerate(header)
                ]
                for row_number, values in enumerate(iterator, start=2):
                    values = list(values)
                    record = {
                        column: _cell_text(values[index] if index < len(values) else None)
                        for index, column in enumerate(columns)
                    }
                    if any(value.strip() for value in record.values()):
                        records.append(record)
                        workbook_values.extend(record.values())
                        for column, value in record.items():
                            if any(marker in value.casefold() for marker in PRODUCTION_BLOCKERS):
                                blocker_locations.append(
                                    f"{file_name}/{worksheet.title}: row {row_number}, column {column}"
                                )

            if worksheet.title in required_sheets and not records:
                raise ValueError(f"Required sheet {file_name}/{worksheet.title} has no data rows.")

            records_json = json.dumps(
                records,
                ensure_ascii=False,
                separators=(",", ":"),
            )
            snapshot.append(
                (file_name, worksheet.title, records_json, source_sha256, loaded_at)
            )
            row_count += len(records)

        workbook.close()

    try:
        validate_production_configuration(catalog, workbook_values)
    except ValueError as error:
        locations = "; ".join(blocker_locations[:5])
        raise ValueError(
            f"{error} Source directory: {root}. "
            f"Found {len(blocker_locations)} blocked cells. First locations: {locations}"
        ) from error
    if not snapshot:
        raise ValueError("No configuration rows were found in the supplied workbooks.")
    return snapshot, row_count


def main(catalog: str, config_dir: str | Path | None = None) -> None:
    if catalog not in ALLOWED_CATALOGS:
        raise ValueError(f"Unsupported Master Data catalog: {catalog}")

    if config_dir:
        source_dir = Path(config_dir)
    else:
        script_path = globals().get("__file__")
        if not script_path:
            raise ValueError(
                "A configuration directory must be supplied when setup_tables.py "
                "runs in notebook context."
            )
        source_dir = Path(script_path).resolve().parent / "data"
    configuration_rows, source_row_count = load_configuration_snapshot(source_dir, catalog)

    spark = SparkSession.builder.getOrCreate()
    requests_table = f"{catalog}.bronze.master_data_app_requests"
    history_table = f"{catalog}.bronze.master_data_app_request_history"
    configuration_table = f"{catalog}.bronze.master_data_app_configuration"

    spark.sql(
        f"""
        CREATE TABLE IF NOT EXISTS {requests_table} (
            request_id STRING,
            request_type STRING,
            customer_id STRING,
            customer_name STRING,
            company_code STRING,
            sales_org STRING,
            status STRING,
            assigned_to STRING,
            requested_by STRING,
            requested_at TIMESTAMP,
            updated_at TIMESTAMP,
            business_justification STRING,
            payload_json STRING,
            original_defaults_json STRING
        ) USING DELTA
        COMMENT 'Customer master data requests submitted through the Master Data App'
        """
    )
    spark.sql(
        f"""
        CREATE TABLE IF NOT EXISTS {history_table} (
            history_id STRING,
            request_id STRING,
            event_type STRING,
            from_status STRING,
            to_status STRING,
            actor STRING,
            comment STRING,
            created_at TIMESTAMP
        ) USING DELTA
        COMMENT 'Status and activity history for Master Data App requests'
        """
    )
    spark.sql(
        f"""
        CREATE TABLE IF NOT EXISTS {configuration_table} (
            source_file STRING,
            sheet_name STRING,
            records_json STRING,
            source_sha256 STRING,
            loaded_at TIMESTAMP
        ) USING DELTA
        COMMENT 'Versioned source workbook snapshots used by the Master Data App'
        """
    )
    configuration_df = spark.createDataFrame(
        configuration_rows,
        schema=CONFIGURATION_TABLE_SCHEMA,
    )
    (
        configuration_df.write.format("delta")
        .mode("overwrite")
        .option("overwriteSchema", "true")
        .saveAsTable(configuration_table)
    )
    print(
        f"Ensured request tables and refreshed {len(configuration_rows)} configuration "
        f"sheets ({source_row_count} source rows) in {catalog}.bronze"
    )


if __name__ == "__main__":
    if len(sys.argv) not in {2, 3}:
        raise SystemExit("Usage: setup_tables.py <automation_catalog> [configuration_directory]")
    main(sys.argv[1], sys.argv[2] if len(sys.argv) == 3 else None)