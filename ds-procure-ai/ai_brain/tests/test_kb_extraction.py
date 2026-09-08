"""Deep extraction: stem bids, chunk catalog, coverage, BidProfile overlay (no live LLM)."""

from __future__ import annotations

import json
import shutil
import tempfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from langgraph.types import Send
from openpyxl import load_workbook

from ai_brain.core.kb.schemas import EvidenceItem, VendorEvidenceBundle
from shared.db.models import utc_now_iso


@pytest.fixture()
def db_env(monkeypatch):
    from shared.db import init_db
    from shared.db import connection as db_connection

    if db_connection._engine is not None:
        db_connection._engine.dispose()
    db_connection._engine = None
    db_connection._SessionLocal = None
    tmp = tempfile.mkdtemp()
    db_path = Path(tmp) / "test.sqlite"
    monkeypatch.setenv("MATE_SQLITE_PATH", str(db_path))
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(Path(tmp) / "projects"))
    init_db()
    yield db_path
    if db_connection._engine is not None:
        db_connection._engine.dispose()
    db_connection._engine = None
    db_connection._SessionLocal = None
    shutil.rmtree(tmp, ignore_errors=True)


def _seed_project(*, extra_artifacts: list[dict] | None = None):
    from shared.db.connection import session_scope
    from shared.db.models import AppUser, Artifact, Project

    now = utc_now_iso()
    with session_scope() as session:
        session.add(AppUser(workspace_id="u1", created_at=now, updated_at=now))
        session.add(
            Project(
                id="p1",
                code="PRJ-TEST",
                owner_id="u1",
                module_id="procurement-ai-assistant",
                name="SWIFT CSP assessment",
                business_process="Indirect",
                category="SWIFT CSCF",
                status="Active",
                priority="Medium",
                progress=0,
                created_at=now,
                updated_at=now,
            )
        )
        for payload in extra_artifacts or []:
            session.add(Artifact(**payload))


def _add_chunk(
    session,
    *,
    chunk_id: str,
    artifact_id: str,
    ordinal: int,
    text: str,
    location: dict,
    heading: str = "Pricing",
):
    from shared.db.models import Chunk

    session.add(
        Chunk(
            id=chunk_id,
            project_id="p1",
            artifact_id=artifact_id,
            ordinal=ordinal,
            chunk_type="text",
            text=text,
            token_count=len(text.split()),
            heading_path_json=json.dumps([heading]),
            location_json=json.dumps(location),
            created_at=utc_now_iso(),
        )
    )


def test_classify_from_stem_unknown_filename():
    from ai_brain.core.kb.classify import STEM_FALLBACK_CONFIDENCE, classify_from_stem

    result = classify_from_stem("Acme_Offer.pdf")
    assert result is not None
    assert result.bid_id == "acme-offer"
    assert result.doc_role == "proposal"
    assert result.confidence == STEM_FALLBACK_CONFIDENCE


def test_classify_and_pin_stem_when_llm_fails(db_env, monkeypatch):
    from ai_brain.core.kb.classify import classify_and_pin, classify_from_stem
    from ai_brain.core.kb.materialize import list_bid_ids

    monkeypatch.setattr(
        "ai_brain.core.kb.classify.classify_snippets",
        lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("no llm")),
    )
    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
            }
        ]
    )
    assert classify_from_stem("Acme_Offer.pdf").bid_id == "acme-offer"
    out = classify_and_pin("p1", "art-acme", "Acme_Offer.pdf")
    assert out["vendor_id"] == "acme-offer"
    assert "acme-offer" in list_bid_ids("p1")


def test_plan_kb_build_fans_out_for_stem_bids(db_env, monkeypatch):
    from ai_brain.core.kb.pipeline import plan_kb_build

    monkeypatch.setattr(
        "ai_brain.core.kb.classify.classify_snippets",
        lambda *_a, **_k: (_ for _ in ()).throw(RuntimeError("no llm")),
    )
    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "parsed_json": json.dumps({"markdown": "Acme year-1 fee 10000 EUR."}),
            },
            {
                "id": "art-rfp",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Etex_RFP_Swift CSP.pdf",
                "storage_relpath": "rel/rfp",
                "created_at": now,
                "parse_status": "ok",
                "doc_role": "rfp",
            },
        ]
    )
    out = plan_kb_build({"project_id": "p1"})
    assert isinstance(out, list)
    assert all(isinstance(item, Send) for item in out)
    assert {item.arg["bid_id"] for item in out} == {"acme-offer"}


def test_list_bid_chunks_returns_every_seeded_chunk_including_sheet(db_env):
    from ai_brain.core.kb.mcp_tools import list_bid_chunk_ids, list_bid_chunks
    from ai_brain.core.kb.search import ChunkHit, _locator_from_location, render_with_provenance
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.xlsx",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-page",
            artifact_id="art-acme",
            ordinal=0,
            text="Acme delivers SWIFT CSP questionnaires.",
            location={"page": 2},
            heading="Approach",
        )
        _add_chunk(
            session,
            chunk_id="c-sheet",
            artifact_id="art-acme",
            ordinal=1,
            text="Day rate 900 EUR on Rate card.",
            location={"sheet": "Rate card"},
            heading="Rates",
        )
    ids = list_bid_chunk_ids("p1", "acme-offer")
    assert ids == ["c-page", "c-sheet"]
    manifest = list_bid_chunks("acme-offer", project_id="p1")
    assert "c-page" in manifest and "c-sheet" in manifest
    assert "sheet Rate card" in manifest
    assert _locator_from_location(json.dumps({"sheet": "Rate card"})) == "sheet Rate card"
    rendered = render_with_provenance(
        [
            ChunkHit(
                chunk_id="c-sheet",
                artifact_id="art-acme",
                original_name="Acme_Offer.xlsx",
                page=None,
                heading="Rates",
                ordinal=1,
                text="Day rate 900 EUR on Rate card.",
                locator="sheet Rate card",
            )
        ]
    )
    assert "Acme_Offer.xlsx" in rendered
    assert "sheet Rate card" in rendered
    assert "Day rate 900" in rendered


def test_read_chunks_records_coverage_and_pages_at_cap(db_env):
    from ai_brain.core.kb.coverage import CoverageState, coverage_scope
    from ai_brain.core.kb.mcp_tools import read_chunk, read_chunks
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
            }
        ]
    )
    big = "rate " * 2500
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c1",
            artifact_id="art-acme",
            ordinal=0,
            text=big,
            location={"page": 1},
        )
        _add_chunk(
            session,
            chunk_id="c2",
            artifact_id="art-acme",
            ordinal=1,
            text=big,
            location={"page": 2},
        )
        _add_chunk(
            session,
            chunk_id="c3",
            artifact_id="art-acme",
            ordinal=2,
            text=big,
            location={"page": 3},
        )
    state = CoverageState(required={"c1", "c2", "c3"})
    with coverage_scope(state):
        body = read_chunks(["c1", "c2", "c3"])
        read_chunk("c1")
    assert "c1" in state.read
    assert "remaining unread" in body
    assert "c3" in body
    assert "c3" not in state.read


def test_coverage_middleware_counts_read_chunk_ids():
    from ai_brain.core.kb.coverage import CoverageMiddleware, CoverageState, coverage_scope, record_tool_request

    state = CoverageState(required={"c1", "c2"})
    mw = CoverageMiddleware()
    with coverage_scope(state):
        req = SimpleNamespace(tool_call={"name": "read_chunk", "args": {"chunk_id": "c1"}})
        assert mw.wrap_tool_call(req, lambda _r: "ok") == "ok"
        record_tool_request(
            SimpleNamespace(tool_call={"name": "read_chunks", "args": {"chunk_ids": ["c2"]}})
        )
    assert state.read == {"c1", "c2"}
    assert state.unread() == []


@pytest.mark.asyncio
async def test_invoke_vendor_analyst_prefetch_skips_coverage_refill(db_env, monkeypatch):
    from contextlib import asynccontextmanager

    from langchain_core.messages import AIMessage

    from ai_brain.core.kb import agents, pipeline
    from ai_brain.core.kb import mcp_tools
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
                "parsed_json": json.dumps({"markdown": "Acme year-1 fee."}),
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-unread",
            artifact_id="art-acme",
            ordinal=0,
            text="Unread commercial terms.",
            location={"page": 1},
        )

    calls: list[dict] = []

    @asynccontextmanager
    async def fake_tools():
        yield []

    monkeypatch.setattr(mcp_tools, "mcp_kb_tools_context", fake_tools)

    class _Research:
        async def ainvoke(self, payload, config=None):
            calls.append(payload)
            return {"messages": list(payload.get("messages") or []) + [AIMessage(content="noted")]}

    class _Emit:
        async def ainvoke(self, payload, config=None):
            return {
                "structured_response": VendorEvidenceBundle(bid_id="acme-offer", items=[]),
                "messages": list(payload.get("messages") or []),
            }

    monkeypatch.setattr(agents, "build_vendor_analyst", lambda **_k: _Research())
    monkeypatch.setattr(agents, "build_vendor_emit", lambda **_k: _Emit())
    monkeypatch.setattr(agents, "build_pricing_extractor", lambda **_k: _Emit())
    monkeypatch.setattr(agents, "build_profile_extractor", lambda **_k: _Emit())

    bundle, _pricing, _n, _profile = await pipeline.invoke_vendor_analyst("p1", "acme-offer")
    assert len(calls) == 1
    assert not any("unread chunks" in note and "c-unread" in note for note in bundle.notes)


def test_bid_invoke_payload_indexes_chunks_without_pasting_files(db_env):
    from ai_brain.core.kb.agents import bid_invoke_payload
    from shared.db.connection import session_scope

    now = utc_now_iso()
    markdown = "Acme unique corpus token ZXCV-9988 appears only in the file body."
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
                "parsed_json": json.dumps({"markdown": markdown}),
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-idx",
            artifact_id="art-acme",
            ordinal=0,
            text=markdown,
            location={"page": 1},
        )
    payload = bid_invoke_payload("p1", "acme-offer")
    content = payload["messages"][0]["content"]
    assert "c-idx" in content
    assert "ZXCV-9988" not in content
    assert any("/vendors/acme-offer/" in path for path in payload["files"])
    assert len(content) < 8_000
    packed = payload.get("packed_chunk_ids") or []
    assert packed == ["c-idx"]
    chunk_files = [path for path in payload["files"] if "/_chunks/" in path]
    assert any(path.endswith("/000.md") for path in chunk_files)
    packed_text = " ".join(
        str((payload["files"][path] or {}).get("content") or payload["files"][path])
        for path in chunk_files
    )
    assert "c-idx" in packed_text
    assert "ZXCV-9988" in packed_text


def test_prefetch_marks_coverage_read_before_research(db_env):
    from ai_brain.core.kb.agents import bid_invoke_payload
    from ai_brain.core.kb.coverage import CoverageState
    from ai_brain.core.kb.mcp_tools import list_bid_chunk_ids, prefetch_bid_chunks
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-a",
            artifact_id="art-acme",
            ordinal=0,
            text="First packed chunk.",
            location={"page": 1},
        )
        _add_chunk(
            session,
            chunk_id="c-b",
            artifact_id="art-acme",
            ordinal=1,
            text="Second packed chunk.",
            location={"page": 2},
        )
    files, packed = prefetch_bid_chunks("p1", "acme-offer")
    assert packed == ["c-a", "c-b"]
    assert any(path.endswith("/_chunks/000.md") for path in files)
    packed_body = "\n".join(files.values())
    assert "`c-a`" in packed_body and "`c-b`" in packed_body
    payload = bid_invoke_payload("p1", "acme-offer")
    coverage = CoverageState(required=set(list_bid_chunk_ids("p1", "acme-offer")))
    coverage.read.update(payload["packed_chunk_ids"])
    assert coverage.unread() == []


def test_read_chunk_window_pages_by_offset(db_env):
    from ai_brain.core.kb.coverage import CoverageState, coverage_scope
    from ai_brain.core.kb.mcp_tools import read_chunk_window
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
            }
        ]
    )
    with session_scope() as session:
        for index in range(3):
            _add_chunk(
                session,
                chunk_id=f"c-{index}",
                artifact_id="art-acme",
                ordinal=index,
                text=f"Window body {index}.",
                location={"page": index + 1},
            )
    state = CoverageState(required={"c-0", "c-1", "c-2"})
    with coverage_scope(state):
        first = read_chunk_window("acme-offer", offset=0, limit=2, project_id="p1")
        second = read_chunk_window("acme-offer", offset=2, limit=2, project_id="p1")
    assert "[c-0]" in first and "[c-1]" in first
    assert "next offset=2" in first
    assert "[c-2]" in second
    assert state.read == {"c-0", "c-1", "c-2"}


def test_filesystem_middleware_allowlist_hides_writes():
    from deepagents.backends import StateBackend
    from deepagents.middleware.filesystem import FilesystemMiddleware

    mw = FilesystemMiddleware(
        backend=StateBackend(),
        tools=["read_file", "ls", "glob", "grep"],
    )
    names = {getattr(tool, "name", "") for tool in mw.tools}
    assert "read_file" in names
    assert "write_file" not in names
    assert "edit_file" not in names


def test_chunk_id_only_citation_grounds_against_sql_chunk(db_env):
    from sqlalchemy import select

    from ai_brain.core.kb.persist import upsert_evidence_bundle
    from ai_brain.core.kb.schemas import apply_grounding_gate
    from shared.db.connection import session_scope
    from shared.db.models import Evidence

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-ok",
            artifact_id="art-acme",
            ordinal=0,
            text="Acme will assess CSCF 2026 independently.",
            location={"sheet": "Scope"},
        )
    item = EvidenceItem(
        requirement_id="R01",
        status="met",
        claim_text="Acme will assess CSCF 2026 independently.",
        chunk_id="c-ok",
        confidence=0.9,
    )
    assert apply_grounding_gate(item).status == "met"
    upsert_evidence_bundle(
        "p1",
        VendorEvidenceBundle(bid_id="acme-offer", items=[item]),
    )
    with session_scope() as session:
        rows = list(session.scalars(select(Evidence).where(Evidence.project_id == "p1")).all())
        assert len(rows) == 1
        assert rows[0].status == "met"
        assert rows[0].chunk_id == "c-ok"


def test_upsert_pricing_writes_lines(db_env):
    from sqlalchemy import select

    from ai_brain.core.kb.persist import upsert_pricing
    from ai_brain.core.kb.schemas import PricingExtraction, PricingLineItem
    from shared.db.connection import session_scope
    from shared.db.models import PricingLine

    _seed_project()
    written = upsert_pricing(
        "p1",
        PricingExtraction(
            bid_id="bdo",
            lines=[
                PricingLineItem(
                    offer_type="fixed_price",
                    role="SWIFT CSP Independent Assessment",
                    line_total=48000,
                    currency="EUR",
                )
            ],
        ),
    )
    assert written == 1
    with session_scope() as session:
        rows = list(session.scalars(select(PricingLine).where(PricingLine.project_id == "p1")).all())
        assert len(rows) == 1
        assert rows[0].vendor_id == "bdo"
        assert rows[0].role == "SWIFT CSP Independent Assessment"
        assert float(rows[0].line_total) == 48000


def test_matrix_from_profile_and_evidence_fills_xlsx(db_env):
    from ai_brain.core.kb.persist import seed_swift_requirements, upsert_bid_profile, upsert_evidence_bundle
    from ai_brain.core.kb.schemas import BidProfile
    from ai_brain.core.procurement_ai.insights import build_insight_payload
    from ai_brain.core.procurement_ai.packs.compare_xlsx import build_comparison_xlsx
    from ai_brain.core.procurement_ai.packs.matrix_from_evidence import load_matrix_source, matrix_from_source
    from shared.db.connection import session_scope

    now = utc_now_iso()
    _seed_project(
        extra_artifacts=[
            {
                "id": "art-acme",
                "project_id": "p1",
                "uploaded_by": "u1",
                "kind": "upload",
                "original_name": "Acme_Offer.pdf",
                "storage_relpath": "rel/acme",
                "created_at": now,
                "parse_status": "ok",
                "vendor_id": "acme-offer",
                "doc_role": "proposal",
                "parsed_json": json.dumps({"markdown": "Acme will assess CSCF 2026 independently."}),
            }
        ]
    )
    with session_scope() as session:
        _add_chunk(
            session,
            chunk_id="c-ok",
            artifact_id="art-acme",
            ordinal=0,
            text="Acme will assess CSCF 2026 independently.",
            location={"page": 3},
        )
    seed_swift_requirements("p1")
    upsert_evidence_bundle(
        "p1",
        VendorEvidenceBundle(
            bid_id="acme-offer",
            items=[
                EvidenceItem(
                    requirement_id="R01",
                    status="met",
                    claim_text="Acme will assess CSCF 2026 independently.",
                    chunk_id="c-ok",
                    confidence=0.9,
                )
            ],
        ),
    )
    upsert_bid_profile(
        "p1",
        BidProfile(
            bid_id="acme-offer",
            timeline="12 weeks from kickoff",
            ecovadis="Gold 72",
            y2_tco="EUR 21000",
            delivery_model="On-site Belgium",
        ),
    )
    source = load_matrix_source("p1")
    assert source.warm
    assert source.profiles["acme-offer"]["timeline"] == "12 weeks from kickoff"
    matrix = matrix_from_source(source, "Indirect")
    vendor = matrix.vendors[0]
    assert vendor.timeline == "12 weeks from kickoff"
    assert vendor.ecovadis == "Gold 72"
    assert vendor.y2_tco == "EUR 21000"
    r01 = next(cell for cell in matrix.requirements if cell.requirement_id == "R01")
    assert any("assess CSCF" in mark for mark in r01.vendor_marks.values())

    class _Project:
        id = "p1"
        code = "PRJ-TEST"
        name = "SWIFT CSP assessment"
        business_process = "Indirect"
        category = "SWIFT CSCF"
        description = "Professional services days"

    class _Artifact:
        id = "art-acme"
        original_name = "Acme_Offer.pdf"
        parse_status = "ok"

    xlsx = build_comparison_xlsx(
        "p1",
        build_insight_payload(_Project(), [_Artifact()], []),
        thread_id="t-extract",
    )
    assert xlsx["evidence_used"] is True
    assert xlsx["status"] == "ready"
    wb = load_workbook(xlsx["path"])
    bafo = wb["Final BAFO Comparison"]
    labels = [bafo.cell(row, 1).value for row in range(1, (bafo.max_row or 1) + 1)]
    timeline_row = labels.index("Planning and timeline (weeks) / Delivery time") + 1
    assert any(
        "12 weeks from kickoff" in str(bafo.cell(timeline_row, col).value or "")
        for col in range(1, (bafo.max_column or 1) + 1)
    )
    req = wb["Procurement Requirements"]
    joined = " ".join(
        str(req.cell(row, col).value or "")
        for row in range(1, (req.max_row or 1) + 1)
        for col in range(1, (req.max_column or 1) + 1)
    )
    assert "assess CSCF" in joined


def test_databricks_safe_messages_strip_tool_result_ids():
    from langchain_core.messages import AIMessage, HumanMessage, ToolMessage

    from ai_brain.core.kb.pipeline import _databricks_safe_messages

    raw = [
        HumanMessage(content="Analyse bid `pwc` only."),
        AIMessage(
            content="",
            tool_calls=[{"name": "read_chunk", "id": "call_1", "args": {"chunk_id": "c1"}}],
        ),
        ToolMessage(
            content=[{"type": "text", "text": "PwC year-1 fee 13410 EUR", "id": "msg_abc"}],
            tool_call_id="call_1",
            name="read_chunk",
        ),
    ]
    out = _databricks_safe_messages(raw)
    assert len(out) == 1
    assert out[0]["role"] == "user"
    assert isinstance(out[0]["content"], str)
    assert "13410" in out[0]["content"]
    assert '"id"' not in out[0]["content"]
    assert "msg_abc" not in out[0]["content"]
    dumped = json.dumps(out)
    assert "tool_result" not in dumped
    assert "tool_calls" not in dumped


@pytest.mark.asyncio
async def test_invoke_vendor_analyst_emit_gets_plain_transcript(db_env, monkeypatch):
    from contextlib import asynccontextmanager

    from langchain_core.messages import AIMessage, HumanMessage, ToolMessage

    from ai_brain.core.kb import agents, pipeline
    from ai_brain.core.kb import mcp_tools

    _seed_project()
    seen: dict[str, list] = {}

    @asynccontextmanager
    async def fake_tools():
        yield []

    monkeypatch.setattr(mcp_tools, "mcp_kb_tools_context", fake_tools)

    class _Research:
        async def ainvoke(self, _payload, config=None):
            return {
                "messages": [
                    HumanMessage(content="analyse pwc"),
                    AIMessage(
                        content="",
                        tool_calls=[
                            {"name": "read_chunk", "id": "call_1", "args": {"chunk_id": "c1"}}
                        ],
                    ),
                    ToolMessage(
                        content=[
                            {
                                "type": "text",
                                "text": "PwC year-1 fee 13410",
                                "id": "msg_abc",
                            }
                        ],
                        tool_call_id="call_1",
                        name="read_chunk",
                    ),
                ]
            }

    class _Emit:
        async def ainvoke(self, payload, config=None):
            seen["messages"] = list(payload.get("messages") or [])
            return {
                "structured_response": VendorEvidenceBundle(bid_id="kpmg", items=[]),
                "messages": payload.get("messages") or [],
            }

    monkeypatch.setattr(agents, "build_vendor_analyst", lambda **_k: _Research())
    monkeypatch.setattr(agents, "build_vendor_emit", lambda **_k: _Emit())
    monkeypatch.setattr(agents, "build_pricing_extractor", lambda **_k: _Emit())
    monkeypatch.setattr(agents, "build_profile_extractor", lambda **_k: _Emit())
    bundle, _pricing, _n, _profile = await pipeline.invoke_vendor_analyst("p1", "kpmg")
    assert bundle.bid_id == "kpmg"
    assert seen["messages"]
    for msg in seen["messages"]:
        role = pipeline._message_role(msg)
        assert role not in {"tool", "function", "ai", "assistant"}
        content = msg.get("content") if isinstance(msg, dict) else getattr(msg, "content", None)
        assert isinstance(content, str)
    blob = " ".join(
        str(msg.get("content") if isinstance(msg, dict) else getattr(msg, "content", ""))
        for msg in seen["messages"]
    )
    assert "13410" in blob
    assert "msg_abc" not in blob


@pytest.mark.asyncio
async def test_invoke_vendor_analyst_emit_pricing_profile_start_in_parallel(
    db_env, monkeypatch
):
    import asyncio
    from contextlib import asynccontextmanager

    from langchain_core.messages import AIMessage

    from ai_brain.core.kb import agents, pipeline
    from ai_brain.core.kb import mcp_tools
    from ai_brain.core.kb.schemas import BidProfile, PricingExtraction

    _seed_project()
    started = 0
    gate = asyncio.Event()
    order: list[str] = []

    @asynccontextmanager
    async def fake_tools():
        yield []

    monkeypatch.setattr(mcp_tools, "mcp_kb_tools_context", fake_tools)

    class _Research:
        async def ainvoke(self, payload, config=None):
            return {
                "messages": list(payload.get("messages") or [])
                + [AIMessage(content="noted")]
            }

    def _peer(name: str, structured):
        class _Agent:
            async def ainvoke(self, payload, config=None):
                nonlocal started
                order.append(name)
                started += 1
                if started >= 3:
                    gate.set()
                await asyncio.wait_for(gate.wait(), timeout=2)
                return {
                    "structured_response": structured,
                    "messages": payload.get("messages") or [],
                }

        return _Agent()

    monkeypatch.setattr(agents, "build_vendor_analyst", lambda **_k: _Research())
    monkeypatch.setattr(
        agents,
        "build_vendor_emit",
        lambda **_k: _peer("emit", VendorEvidenceBundle(bid_id="kpmg", items=[])),
    )
    monkeypatch.setattr(
        agents,
        "build_pricing_extractor",
        lambda **_k: _peer("pricing", PricingExtraction(bid_id="kpmg", lines=[])),
    )
    monkeypatch.setattr(
        agents,
        "build_profile_extractor",
        lambda **_k: _peer(
            "profile", BidProfile(bid_id="kpmg", timeline="12 months")
        ),
    )
    bundle, pricing, _n, profile = await pipeline.invoke_vendor_analyst("p1", "kpmg")
    assert set(order) == {"emit", "pricing", "profile"}
    assert bundle.bid_id == "kpmg"
    assert pricing is not None and pricing.bid_id == "kpmg"
    assert profile is not None and profile.timeline == "12 months"
