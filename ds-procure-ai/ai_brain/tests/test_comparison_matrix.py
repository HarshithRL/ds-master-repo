"""Request-time comparison matrix: ontology, writer richness, MCP tools, parse-on-ask."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from ai_brain.core.procurement_ai.comparison_template import TEMPLATE_ROWS
from ai_brain.core.procurement_ai.ontology import (
    ComparisonMatrix,
    ONTOLOGY_VERSION,
    Requirement,
    VendorOffering,
    VendorProfile,
)
from ai_brain.core.procurement_ai.packs.compare_xlsx import build_comparison_xlsx
from ai_brain.core.procurement_ai.packs.matrix_layout import SWIFT_REQUIREMENTS


def _matrix(**overrides) -> ComparisonMatrix:
    payload = {
        "process_type": "indirect_it_services",
        "ontology_version": ONTOLOGY_VERSION,
        "requirements": [
            Requirement(
                id="R01",
                text="Two BIC codes in scope",
                category="scope",
                target="ETEXBEBB (BE) + URSNESMMXXX (ES)",
            ),
            Requirement(
                id="R18",
                text="Dual pricing: T&M AND Fixed Price",
                category="commercial",
                target="Two offers required",
            ),
        ],
        "vendors": [
            VendorProfile(
                vendor_id="bdo",
                legal_name="BDO Advisory BV",
                y1_service="35251",
                y1_etex_mandays="€3,600 (4.5 days)",
                y1_tco="38851",
                y1_solution="N/A - advisory service, no software",
                blended_rate="€1,152/day discounted (list €1,280)",
                vendor_mandays="34 days mandatory",
                timeline="8–10 wks; start Aug",
                ecovadis="not provided",
                creditsafe="Graydon report",
                missing_requirements="Fixed-price offer missing (RFP requires both)",
                source_files="bdo-proposal.pdf (21 pp)",
                locators={"y1_service": "p.4 bdo-proposal.pdf"},
            ),
            VendorProfile(
                vendor_id="ey",
                legal_name="EY Bedrijfsrevisoren BV",
                y1_service="23800",
                y1_etex_mandays="€16,000 (20 days)",
                y1_tco="39800",
                timeline="Q3 confirmed",
                ecovadis="Silver FY25",
                creditsafe="Creditsafe EY",
                source_files="EY proposal (vF).pdf (12 pp)",
            ),
        ],
        "offerings": [
            VendorOffering(
                vendor_id="bdo",
                vendor_legal_name="BDO Advisory BV",
                requirement_id="R01",
                value_text="✓ both BICs",
                status="✓",
                state="met",
                locator="p.3 bdo-proposal.pdf",
                quote="ETEXBEBB and URSNESMMXXX in scope",
            ),
            VendorOffering(
                vendor_id="bdo",
                vendor_legal_name="BDO Advisory BV",
                requirement_id="R18",
                value_text="✗ T&M only",
                status="✗",
                state="not_offered",
                locator="p.8 bdo-proposal.pdf",
                quote="Time and materials only",
            ),
            VendorOffering(
                vendor_id="ey",
                vendor_legal_name="EY Bedrijfsrevisoren BV",
                requirement_id="R01",
                value_text="✓",
                status="✓",
                state="met",
                locator="p.2 EY proposal (vF).pdf",
                quote="Two BIC codes covered",
            ),
        ],
        "sources": [
            {"vendor_id": "bdo", "files": ["bdo-proposal.pdf (21 pp)"]},
            {"vendor_id": "ey", "files": ["EY proposal (vF).pdf (12 pp)"]},
            {"vendor_id": "rfp", "files": ["Etex_RFP_Swift CSP_06022026.pdf (15 pp)"]},
        ],
        "method_notes": "Cost rows use each vendor's quoted figures. Empty is missing, never zero.",
        "rfp_sources": "Etex_RFP_Swift CSP_06022026.pdf (15 pp)",
        "summary": "Draft SWIFT CSP comparison. Humans award.",
        "key_flags": [
            {
                "severity": "high",
                "text": "BDO offered T&M only; RFP requires dual pricing.",
                "vendors": ["bdo"],
            }
        ],
        "bafo_rows": [row["label"] for row in TEMPLATE_ROWS],
    }
    payload.update(overrides)
    return ComparisonMatrix.model_validate(payload)


def test_template_row_labels_match_bafo_skeleton():
    labels = [row["label"] for row in TEMPLATE_ROWS]
    assert labels[0] == "Solution costs (Software, Materials)"
    assert "Total TCO year 1" in labels
    assert "Score Card score" in labels
    assert "Renewal type" in labels
    assert "Supply type location" in labels
    assert "Reference: Customer Names" in labels
    assert labels[-1] == "Missing requirement list"


def test_offering_requires_quote_when_evidenced():
    with pytest.raises(ValidationError, match="quote"):
        VendorOffering(
            vendor_id="bdo",
            vendor_legal_name="BDO",
            requirement_id="R01",
            status="✓",
            locator="p.1",
            quote="",
        )


def test_direct_rejects_swift_overlay_labels():
    with pytest.raises(ValidationError, match="SWIFT"):
        ComparisonMatrix(
            process_type="direct_tg",
            ontology_version=ONTOLOGY_VERSION,
            requirements=[
                Requirement(
                    id="R01",
                    text="Two BIC codes in scope",
                    category="scope",
                )
            ],
            vendors=[VendorProfile(vendor_id="mill", legal_name="Mill A")],
        )


def test_ecovadis_is_not_creditsafe():
    with pytest.raises(ValidationError, match="separate"):
        VendorProfile(
            vendor_id="ey",
            legal_name="EY",
            ecovadis="72",
            creditsafe="72",
        )


def test_rich_workbook_from_ontology(tmp_path, monkeypatch):
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(tmp_path))
    insights = {
        "process_type": "indirect_it_services",
        "project_name": "SWIFT CSP assessment",
        "project_code": "PRJ-TEST",
        "process_label": "IT services",
        "vendors": [],
        "decision": {},
    }
    matrix = _matrix()
    result = build_comparison_xlsx("p1", insights, thread_id="t-rich", matrix=matrix)
    assert result["status"] == "ready"
    assert result["evidence_used"] is True
    assert result["llm_used"] is False
    wb = load_workbook(result["path"])
    for name in ("Final BAFO Comparison", "Procurement Requirements", "Sources & Notes", "Cover"):
        assert name in wb.sheetnames, wb.sheetnames
    assert "Citations" not in wb.sheetnames
    bafo = wb["Final BAFO Comparison"]
    assert bafo["A3"].value == "Solution costs (Software, Materials)"
    assert bafo["A6"].value == "Total TCO year 1"
    assert bafo["A17"].value == "Score Card score"
    assert bafo["C3"].value == "N/A"
    assert "35251" in str(bafo["C4"].value)
    assert "€3,600" in str(bafo["C5"].value)
    assert "To be scored by Etex evaluation team" in str(bafo["C17"].value)
    vendors = [bafo.cell(1, col).value for col in range(3, 8) if bafo.cell(1, col).value]
    assert "BDO Advisory BV" in vendors
    assert "EY Bedrijfsrevisoren BV" in vendors
    reqs = wb["Procurement Requirements"]
    blob = "\n".join(
        str(value or "")
        for row in reqs.iter_rows(values_only=True)
        for value in row
    )
    assert "✓ both BICs" in blob
    assert "✗ T&M only" in blob
    sources = wb["Sources & Notes"]
    joined = " ".join(str(sources.cell(r, 2).value or "") for r in range(1, 16))
    assert "bdo-proposal.pdf" in joined
    assert "Etex_RFP_Swift CSP_06022026.pdf" in joined
    assert r"D:\Work" not in _sheet_text(wb)


def test_services_marks_software_rows_na(tmp_path, monkeypatch):
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(tmp_path))
    insights = {"process_type": "indirect_it_services", "vendors": [{"name": "BDO"}], "decision": {}}
    result = build_comparison_xlsx("p-na", insights, thread_id="t-na")
    bafo = load_workbook(result["path"])["Final BAFO Comparison"]
    assert bafo["C3"].value == "N/A"
    assert bafo["A17"].value == "Score Card score"


def test_direct_keeps_solution_rows(tmp_path, monkeypatch):
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(tmp_path))
    insights = {
        "process_type": "direct_tg",
        "project_name": "TG screws",
        "vendors": [{"name": "Mill A"}],
        "decision": {},
    }
    result = build_comparison_xlsx("p-direct", insights, thread_id="t-d")
    wb = load_workbook(result["path"])
    assert "extra-polation " in wb.sheetnames
    bafo = wb["Final BAFO Comparison"]
    assert bafo["C3"].value != "N/A"
    assert bafo["C4"].value == "N/A"
    reqs = wb["Procurement Requirements"]
    assert "SWIFT" in str(reqs["A2"].value or "") or "Direct TG" in str(reqs["A2"].value or "")
    assert reqs["A5"].value != "R01" or "Two BIC" not in str(reqs["B5"].value or "")


def test_orchestrator_prompt_lists_bids_and_swift_overlay():
    import inspect

    from ai_brain.core.procurement_ai.subagents.comparison_matrix import (
        build_comparison_orchestrator,
        comparison_user_message,
    )

    msg = comparison_user_message("p1", "indirect_it_services", ["bdo", "ey"])
    assert "bdo" in msg
    assert "R01" in msg
    direct = comparison_user_message("p2", "direct_tg", ["mill-a"])
    assert "SWIFT CSP overlay" not in direct
    assert "Do not paste SWIFT" in direct
    src = inspect.getsource(build_comparison_orchestrator)
    assert 'skills=["/skills/procurement/"]' in src
    assert "CoverageMiddleware" in src


def test_mcp_adapter_available_for_kb_tools():
    from ai_brain.core.kb.api.mcp_tools import _mcp_adapter_class

    cls = _mcp_adapter_class()
    if cls is None:
        pytest.skip("langchain.mcp not installed")
    assert cls.__name__ == "MCPAdapter"


def test_langchain_kb_tools_include_list_bids():
    from ai_brain.core.kb.api.mcp_tools import langchain_kb_tools

    names = {getattr(t, "__name__", str(t)) for t in langchain_kb_tools()}
    assert "list_bids" in names


def test_kb_build_defaults_skip_xlsx():
    import inspect

    from ai_brain.core.kb.extract.runner import arun_kb_build

    assert inspect.signature(arun_kb_build).parameters["write_xlsx"].default is False


def test_compare_xlsx_parses_then_extracts(tmp_path, monkeypatch):
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(tmp_path))
    from ai_brain.core.procurement_ai.subagents import compare_xlsx as compare_mod

    parsed = {"called": False}
    extracted = {"called": False}

    async def fake_parse(project_id):
        parsed["called"] = True
        return None

    async def fake_extract(project_id, insights, config=None):
        extracted["called"] = True
        return _matrix()

    class _Project:
        id = "proj-demo"
        code = "PR-100"
        name = "SWIFT CSP assessment"
        business_process = "Indirect"
        category = "SWIFT CSCF"
        description = "services"

    class _Artifact:
        id = "a1"
        original_name = "BDO.pdf"
        parse_status = "ok"

    monkeypatch.setattr(compare_mod, "_ensure_corpus", fake_parse)
    monkeypatch.setattr(compare_mod, "load_context", lambda _pid: (_Project(), [_Artifact()], []))
    monkeypatch.setattr(compare_mod, "build_insight_payload", lambda *a, **k: {
        "process_type": "indirect_it_services",
        "project_name": "SWIFT CSP assessment",
        "project_code": "PR-100",
        "vendors": [],
        "decision": {},
    })
    monkeypatch.setattr(
        "ai_brain.core.kb.build.pipeline.ensure_vendor_bids",
        lambda _pid: ["bdo"],
    )
    monkeypatch.setattr(
        "ai_brain.core.procurement_ai.subagents.comparison_matrix.extract_comparison_matrix",
        fake_extract,
    )

    out = asyncio.run(compare_mod.compare_xlsx({"project_id": "proj-demo", "procurement": {}}))
    assert parsed["called"] is True
    assert extracted["called"] is True
    assert out["procurement"]["xlsx_status"] == "ready"
    assert out["procurement"]["evidence_used"] is True
    assert Path(tmp_path).exists()


def _sheet_text(wb) -> str:
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    parts.append(str(value))
    return "\n".join(parts)


def test_failed_extract_writes_template_not_fixture(tmp_path, monkeypatch):
    monkeypatch.setenv("MATE_PROJECTS_DATA_ROOT", str(tmp_path))
    insights = {
        "process_type": "indirect_it_services",
        "vendors": [{"name": "BDO"}],
        "decision": {},
    }
    result = build_comparison_xlsx("p-empty", insights, thread_id="t-empty")
    assert result["status"] == "template"
    assert result["evidence_used"] is False
    bafo = load_workbook(result["path"])["Final BAFO Comparison"]
    assert str(bafo["C4"].value) in {"?", "missing", "N/A"}
    assert "35251" not in str(bafo["C4"].value)


def test_swift_overlay_has_forty_requirements():
    assert len(SWIFT_REQUIREMENTS) == 40
    assert SWIFT_REQUIREMENTS[0][0] == "R01"
    assert SWIFT_REQUIREMENTS[-1][0] == "R40"
