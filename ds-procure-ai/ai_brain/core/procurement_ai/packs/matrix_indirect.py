"""Indirect workbook: SWIFT CSP Vendor Comparison.xlsx layout."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ai_brain.core.procurement_ai.comparison_template import (
    FIELD_TO_ATTR,
    NA,
    TEMPLATE_ROWS,
    row_applies,
)
from ai_brain.core.procurement_ai.packs.matrix_layout import (
    MISSING,
    NOT_STATED,
    REQUIREMENTS_LEGEND,
    SCORED,
    SCORED_BY_ETEX,
    STATUS_MARK,
    SWIFT_REQUIREMENTS,
)
from ai_brain.core.procurement_ai.packs.matrix_schema import (
    IndirectMatrix,
    IndirectVendorFacts,
    RequirementCell,
    blank_indirect_vendor,
)
from ai_brain.core.procurement_ai.packs.xlsx_style import (
    NA_FILL,
    NA_FONT,
    SECTION_FILL,
    SECTION_FONT,
    WRAP,
    autosize,
    header_row,
    put,
    put_render,
    vendor_names,
)

LEAD_HEADERS = ("ID", "Requirement", "Category", "Mandatory", "Target / Detail")


def write_indirect(wb, insights: dict[str, Any], matrix: IndirectMatrix) -> None:
    vendors = _vendors(insights, matrix)
    process = str(insights.get("process_type") or "").strip()
    _write_bafo(
        wb.create_sheet("Final BAFO Comparison"),
        vendors,
        process_type=process,
        from_evidence=matrix.from_evidence,
    )
    _write_requirements(wb.create_sheet("Procurement Requirements"), insights, matrix, vendors)
    _write_sources(wb.create_sheet("Sources & Notes"), insights, matrix, vendors)


def _vendors(insights: dict[str, Any], matrix: IndirectMatrix) -> list[IndirectVendorFacts]:
    extracted = {_norm_name(v.name): v for v in matrix.vendors if v.name}
    ordered: list[IndirectVendorFacts] = []
    used: set[str] = set()
    for name in vendor_names(insights, None):
        key = _norm_name(name)
        match = extracted.get(key) or _fuzzy_vendor(extracted, name)
        if match is not None:
            used.add(_norm_name(match.name))
            used.add(key)
            ordered.append(match.model_copy(update={"name": name}))
        else:
            ordered.append(blank_indirect_vendor(name))
    for vendor in matrix.vendors:
        if vendor.name and _norm_name(vendor.name) not in used:
            ordered.append(vendor)
    return ordered or [blank_indirect_vendor("Unnamed vendor")]


def _norm_name(name: str) -> str:
    return " ".join(name.lower().replace("+", " ").split())


def _fuzzy_vendor(extracted: dict[str, IndirectVendorFacts], name: str) -> IndirectVendorFacts | None:
    needle = _norm_name(name)
    for key, vendor in extracted.items():
        if needle in key or key in needle:
            return vendor
    return None


def _write_bafo(
    ws: Worksheet,
    vendors: list[IndirectVendorFacts],
    *,
    process_type: str = "",
    from_evidence: bool = False,
) -> None:
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws["B1"] = "Type "
    ws["B1"].font = Font(bold=True)
    ws["B1"].alignment = Alignment(wrap_text=True, vertical="center")
    for index, vendor in enumerate(vendors):
        col = 3 + index
        ws.cell(1, col, vendor.name).font = Font(bold=True)
        legal = vendor.legal_name or vendor.name
        ws.cell(2, col, f"Vendor {index + 1} - {legal}")
        ws.cell(2, col).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 36
    excel_row = 3
    for row in TEMPLATE_ROWS:
        if row["field_id"] == "scorecard":
            excel_row += 1
        label = row["label"]
        attr = FIELD_TO_ATTR.get(row["field_id"])
        applies = row_applies(row, process_type)
        ws.cell(excel_row, 1, label).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(excel_row, 2, row.get("cadence") or "")
        ws.cell(excel_row, 2).alignment = Alignment(wrap_text=True, vertical="center")
        for index, vendor in enumerate(vendors):
            comment = (vendor.cell_comments or {}).get(attr or "") if attr else ""
            if not applies:
                cell = ws.cell(excel_row, 3 + index, NA)
                cell.fill = NA_FILL
                cell.font = NA_FONT
                cell.alignment = WRAP
                continue
            raw = getattr(vendor, attr, MISSING) if attr else ""
            if label in SCORED_BY_ETEX and (not raw or raw == MISSING):
                put(ws, excel_row, 3 + index, SCORED, comment=comment)
                continue
            state = None
            if from_evidence and (not raw or raw == MISSING):
                raw = f"{STATUS_MARK['not_stated']} {NOT_STATED}"
                state = "not_stated"
            if from_evidence:
                put_render(
                    ws,
                    excel_row,
                    3 + index,
                    raw or MISSING,
                    state=state,
                    comment=comment,
                )
            else:
                put(ws, excel_row, 3 + index, raw or MISSING, comment=comment)
        ws.row_dimensions[excel_row].height = 32 if len(label) > 40 else 22
        excel_row += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    for index in range(len(vendors)):
        ws.column_dimensions[get_column_letter(3 + index)].width = 28
    ws.freeze_panes = "C3"


def _write_requirements(
    ws: Worksheet,
    insights: dict[str, Any],
    matrix: IndirectMatrix,
    vendors: list[IndirectVendorFacts],
) -> None:
    from ai_brain.core.procurement_ai.process_type import DIRECT_TG, IT_SERVICES

    process = str(insights.get("process_type") or "").strip()
    title = (
        f"Procurement Requirements Compliance Matrix - "
        f"{insights.get('project_name') or 'Etex'} "
        f"({insights.get('project_code') or ''})".strip()
    )
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    if process == DIRECT_TG and not matrix.requirements:
        ws["A2"] = "Direct TG — requirements come from this RFP, not the SWIFT CSP R01–R40 overlay."
        return
    if matrix.from_evidence:
        ws["A2"] = REQUIREMENTS_LEGEND
        headers = [*LEAD_HEADERS, *[v.name for v in vendors]]
        header_row(ws, headers, row=4)
        row = 5
        prev_category = None
        last_col = len(headers)
        extracted = [item for item in matrix.requirements if item.label or item.vendor_states]
        if process == IT_SERVICES:
            items = _overlay_requirements(matrix)
        elif process == DIRECT_TG:
            items = extracted
        else:
            items = extracted or _overlay_requirements(matrix)
        for item in items:
            category = (item.category or "").strip()
            if category and category != prev_category:
                _write_section_row(ws, row, last_col, category)
                row += 1
                prev_category = category
            ws.cell(row, 1, item.requirement_id)
            ws.cell(row, 2, item.label).alignment = WRAP
            ws.cell(row, 3, item.category)
            ws.cell(row, 4, "Yes" if item.mandatory else "No")
            ws.cell(row, 5, item.target)
            for index, vendor in enumerate(vendors):
                mark = (item.vendor_marks or {}).get(vendor.name)
                state = (item.vendor_states or {}).get(vendor.name)
                if not mark:
                    mark = f"{STATUS_MARK['not_stated']} {NOT_STATED}"
                    state = "not_stated"
                pinned = bool((item.vendor_pinned or {}).get(vendor.name)) or "[reviewed]" in str(mark)
                comment = (item.vendor_comments or {}).get(vendor.name)
                citation = (item.vendor_citations or {}).get(vendor.name)
                if citation and not comment:
                    comment = citation
                put_render(
                    ws,
                    row,
                    6 + index,
                    mark,
                    state=state,
                    pinned=pinned,
                    comment=comment,
                )
            row += 1
        ws.freeze_panes = "F5"
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 36
    else:
        ws["A2"] = REQUIREMENTS_LEGEND
        headers = ["ID", "Requirement", "Category", "Target / Detail", *[v.name for v in vendors]]
        header_row(ws, headers, row=4)
        marks_by_id = {item.requirement_id: item.vendor_marks for item in matrix.requirements}
        for offset, (rid, label, category, target) in enumerate(SWIFT_REQUIREMENTS, start=5):
            ws.cell(offset, 1, rid)
            ws.cell(offset, 2, label)
            ws.cell(offset, 3, category)
            ws.cell(offset, 4, target)
            vendor_marks = marks_by_id.get(rid) or {}
            for index, vendor in enumerate(vendors):
                put(ws, offset, 5 + index, vendor_marks.get(vendor.name) or MISSING)
        ws.freeze_panes = "E5"
        ws.column_dimensions["D"].width = 36
    autosize(ws, max_width=36)
    ws.column_dimensions["B"].width = 48


def _overlay_requirements(matrix: IndirectMatrix) -> list[RequirementCell]:
    """SWIFT R01–R40 spine plus any extra SQL requirements."""
    by_id = {item.requirement_id: item for item in matrix.requirements}
    rows: list[RequirementCell] = []
    used: set[str] = set()
    for rid, label, category, target in SWIFT_REQUIREMENTS:
        used.add(rid)
        item = by_id.get(rid)
        if item is None:
            rows.append(
                RequirementCell(
                    requirement_id=rid,
                    label=label,
                    category=category,
                    target=target,
                    mandatory=True,
                )
            )
        else:
            rows.append(
                item.model_copy(
                    update={
                        "label": item.label or label,
                        "category": item.category or category,
                        "target": item.target or target,
                    }
                )
            )
    for item in matrix.requirements:
        if item.requirement_id not in used:
            rows.append(item)
    return rows


def _write_section_row(ws: Worksheet, row: int, last_col: int, category: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, category)
    cell.fill = SECTION_FILL
    cell.font = SECTION_FONT
    cell.alignment = WRAP
    for col in range(1, last_col + 1):
        ws.cell(row, col).fill = SECTION_FILL
        ws.cell(row, col).font = SECTION_FONT


def _write_sources(
    ws: Worksheet,
    insights: dict[str, Any],
    matrix: IndirectMatrix,
    vendors: list[IndirectVendorFacts],
) -> None:
    ws["A1"] = "Data sources (parsed project uploads)"
    ws["A1"].font = Font(bold=True)
    row = 3
    for index, vendor in enumerate(vendors, start=1):
        ws.cell(row, 1, f"Vendor {index} - {vendor.legal_name or vendor.name}")
        files = vendor.source_files
        if not files:
            insight = next((v for v in insights.get("vendors") or [] if v.get("name") == vendor.name), None)
            files = (insight or {}).get("headline") or MISSING
        put(ws, row, 2, files)
        row += 1
    ws.cell(row, 1, "RFP")
    put(ws, row, 2, matrix.rfp_sources or "See uploaded RFP / requirements files")
    row += 2
    ws.cell(row, 1, "Method notes")
    put(
        ws,
        row,
        2,
        matrix.method_notes
        or "Cost rows use each vendor's quoted figures when present. "
        "Etex man-days valued at EUR 800/day. Empty is missing, never zero. Humans award.",
    )
    row += 1
    ws.cell(row, 1, "Key flags")
    flags = matrix.flags or (insights.get("decision") or {}).get("blockers") or []
    put(ws, row, 2, "; ".join(str(item) for item in flags) or MISSING)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 80
