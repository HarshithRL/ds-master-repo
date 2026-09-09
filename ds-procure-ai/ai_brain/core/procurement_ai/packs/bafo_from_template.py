"""Deterministic Final BAFO workbook from the shared template. No invented numbers."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ai_brain.core.procurement_ai.comparison_template import (
    MISSING,
    NA,
    SCORECARD,
    TEMPLATE_ROWS,
    row_applies,
)
from ai_brain.core.procurement_ai.process_type import process_chip

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NA_FILL = PatternFill("solid", fgColor="F2F2F2")
MISSING_FONT = Font(color="C65911")
NOTE = (
    "Draft. Humans award. N/A means the row does not apply to this process type. "
    "? means no evidence. Never treat ? as 0. Scorecard is not auto-scored. "
    "EcoVadis and CreditScore are different fields."
)


def vendor_columns(insights: dict[str, Any]) -> list[str]:
    names: list[str] = []
    raw = insights.get("vendors") or insights.get("vendor_names") or []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, str) and item.strip():
                names.append(item.strip())
            elif isinstance(item, dict):
                name = str(item.get("name") or item.get("vendor") or "").strip()
                if name:
                    names.append(name)
    # Placeholder column only. Not a real vendor name from another process type.
    return names or ["Vendor 1"]


def _fact(facts: dict[str, Any], vendor: str, field_id: str) -> Any:
    by_vendor = facts.get(vendor) if isinstance(facts.get(vendor), dict) else None
    if by_vendor is None and isinstance(facts.get("cells"), dict):
        by_vendor = facts["cells"].get(vendor)
    if not isinstance(by_vendor, dict):
        return None
    return by_vendor.get(field_id)


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def cell_value(
    row: dict[str, Any],
    process_type: str,
    vendor: str,
    facts: dict[str, Any],
    written: dict[tuple[str, str], Any],
) -> Any:
    if not process_type:
        return MISSING
    if not row_applies(row, process_type):
        return NA
    kind = row.get("kind")
    if kind == "scorecard":
        return SCORECARD
    if kind == "total":
        parts = []
        for field_id in row.get("components") or ():
            value = written.get((vendor, field_id), MISSING)
            if value == NA:
                continue
            if not _is_number(value):
                return MISSING
            parts.append(value)
        if not parts:
            return NA
        return sum(parts)
    evidenced = _fact(facts, vendor, row["field_id"])
    if evidenced is None or evidenced == "":
        return MISSING
    if evidenced == 0 or evidenced == "0":
        # Zero is allowed only when the fact explicitly says evidenced.
        if not (isinstance(facts.get(vendor), dict) and facts[vendor].get(f"{row['field_id']}__evidenced")):
            return MISSING
    return evidenced


def write_bafo_workbook(
    path: Path,
    insights: dict[str, Any],
    facts: dict[str, Any] | None = None,
) -> dict[str, Any]:
    process_type = str(insights.get("process_type") or "").strip()
    vendors = vendor_columns(insights)
    facts = facts or {}
    placeholder = vendors == ["Vendor 1"]
    wb = Workbook()
    bafo = wb.active
    bafo.title = "Final BAFO Comparison"
    _write_bafo(bafo, wb, process_type, vendors, facts, placeholder)
    _write_requirements(wb.create_sheet("Procurement Requirements"), process_type)
    _write_notes(wb.create_sheet("Sources & Notes"), process_type, placeholder, facts)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {
        "path": str(path),
        "process_type": process_type,
        "vendors": vendors,
        "placeholder_vendors": placeholder,
        "sheets": list(wb.sheetnames),
    }


def _write_bafo(ws, wb, process_type: str, vendors: list[str], facts: dict[str, Any], placeholder: bool) -> None:
    ws["A1"] = "Type"
    ws["B1"] = "Cadence"
    for index, name in enumerate(vendors, start=3):
        ws.cell(1, index, name)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws["A2"] = NOTE
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2 + len(vendors))
    written: dict[tuple[str, str], Any] = {}
    excel_row = 3
    for row in TEMPLATE_ROWS:
        ws.cell(excel_row, 1, row["label"]).alignment = Alignment(wrap_text=True)
        ws.cell(excel_row, 2, row.get("cadence") or "")
        for index, vendor in enumerate(vendors, start=3):
            value = cell_value(row, process_type, vendor, facts, written)
            written[(vendor, row["field_id"])] = value
            cell = ws.cell(excel_row, index, value)
            if value == NA:
                cell.fill = NA_FILL
            elif value == MISSING:
                cell.font = MISSING_FONT
        # Named field points at the first vendor cell of this row.
        first = f"'{ws.title}'!${get_column_letter(3)}${excel_row}"
        wb.defined_names.add(DefinedName(name=row["field_id"], attr_text=first))
        excel_row += 1
    if placeholder:
        ws.cell(excel_row + 1, 1, "No vendor evidence. Vendor 1 is a placeholder column, not a named bidder.")
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 22
    for index in range(3, 3 + len(vendors)):
        ws.column_dimensions[get_column_letter(index)].width = 28
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A1:{get_column_letter(2 + len(vendors))}{excel_row - 1}"


def _write_requirements(ws, process_type: str) -> None:
    ws["A1"] = "Requirement"
    ws["B1"] = "Status"
    ws["C1"] = "Locator"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    slots: list[str] = []
    if process_type:
        try:
            from ai_brain.core.procurement_ai.fixture_loader import load_catalog

            catalog = load_catalog(process_type)
            for item in getattr(catalog, "requirements", []) or []:
                label = f"{item.id} {item.text}".strip()
                if label:
                    slots.append(label)
        except Exception:
            slots = []
    if not slots:
        ws["A2"] = "No requirement catalog for this process type yet."
        ws["B2"] = MISSING if process_type else "process_type unset"
        return
    for index, label in enumerate(slots, start=2):
        ws.cell(index, 1, label)
        ws.cell(index, 2, MISSING)
        ws.cell(index, 3, "")
    ws.column_dimensions["A"].width = 64


def _write_notes(ws, process_type: str, placeholder: bool, facts: dict[str, Any]) -> None:
    ws["A1"] = "Sources & Notes"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Process type"
    ws["B3"] = process_chip(process_type) if process_type else "Unset"
    ws["A4"] = "Legend"
    ws["B4"] = "N/A = does not apply. ? = no evidence, never 0. Scorecard is human-scored."
    ws["A5"] = "Rate note"
    ws["B5"] = "Etex mandays use the template rate 800 EUR/day. The count still needs evidence."
    row = 6
    if not process_type:
        ws.cell(row, 1, "Blocker")
        ws.cell(row, 2, "process_type unset. Cells are ?, not N/A, until Direct or Indirect is known.")
        row += 1
    if placeholder:
        ws.cell(row, 1, "Vendors")
        ws.cell(row, 2, "No named vendors in evidence. Placeholder column only.")
        row += 1
    notes = facts.get("notes") if isinstance(facts.get("notes"), list) else []
    if not notes:
        ws.cell(row, 1, "Evidence")
        ws.cell(row, 2, "No sourced notes yet. Do not copy numbers from another process type.")
        return
    for note in notes:
        ws.cell(row, 1, "Note")
        ws.cell(row, 2, str(note))
        row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 96
